"""Integration test: τρέχει το index.html σε browser με συνθετικά δεδομένα,
κάνει download το export, και επαληθεύει ότι οι εγγραφές με TRANS πάνε
στον δεύτερο πίνακα του φύλλου «Πελάτες ανά Περιοχή».
"""
import os
import sys
import time
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

HERE = os.path.dirname(__file__)
ROOT = os.path.abspath(os.path.join(HERE, ".."))
FIX = os.path.join(HERE, "fixtures")
INDEX_URL = "file:///" + os.path.join(ROOT, "index.html").replace("\\", "/")
MASTER_PATH = os.path.join(FIX, "master.xlsx")
INFO_PATH = os.path.join(FIX, "INFO.xlsx")
DOWNLOAD_PATH = os.path.join(FIX, "_downloaded.xlsx")


def run():
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        ctx = browser.new_context(accept_downloads=True)
        page = ctx.new_page()
        logs = []
        page.on("console", lambda msg: logs.append(f"[{msg.type}] {msg.text}"))
        page.on("pageerror", lambda exc: logs.append(f"[pageerror] {exc}"))

        try:
            page.goto(INDEX_URL)
            page.wait_for_selector("#masterInput")
            page.set_input_files("#masterInput", MASTER_PATH)
            page.set_input_files("#infoInput", INFO_PATH)

            # Περίμενε να εμφανιστεί το export button (δημιουργείται async μετά την προβολή)
            page.wait_for_selector("#exportFinalExcelBtn", timeout=60000, state="attached")

            with page.expect_download(timeout=30000) as dl_info:
                page.click("#exportFinalExcelBtn")
            dl = dl_info.value
            if os.path.exists(DOWNLOAD_PATH):
                os.remove(DOWNLOAD_PATH)
            dl.save_as(DOWNLOAD_PATH)
        finally:
            with open(os.path.join(FIX, "_console.log"), "w", encoding="utf-8") as f:
                f.write("\n".join(logs))
            try:
                page.screenshot(path=os.path.join(FIX, "_screenshot.png"), full_page=True)
            except Exception:
                pass
            browser.close()

    # === Verify ===
    wb = load_workbook(DOWNLOAD_PATH, data_only=True)
    print("Sheets:", wb.sheetnames)
    assert "Πελάτες ανά Περιοχή" in wb.sheetnames, "Λείπει το νέο φύλλο"
    ws = wb["Πελάτες ανά Περιοχή"]

    # Διάβασε όλα τα cells
    rows = []
    for r in range(1, ws.max_row + 1):
        rows.append([ws.cell(r, c).value for c in range(1, ws.max_column + 1)])

    # Βρες τις γραμμές τίτλων (μ' ένα κελί merged που γράφει "εντός Αττικής" / "με Μεταφορική")
    intra_title_row = None
    trans_title_row = None
    for i, row in enumerate(rows, start=1):
        for cell in row:
            s = str(cell or "")
            if "εντός Αττικής" in s and intra_title_row is None:
                intra_title_row = i
            if "Μεταφορικ" in s and trans_title_row is None:
                trans_title_row = i

    print("intra_title_row=", intra_title_row, "trans_title_row=", trans_title_row)
    assert intra_title_row, "Δεν βρέθηκε τίτλος intra"
    assert trans_title_row, "Δεν βρέθηκε τίτλος trans"
    assert trans_title_row > intra_title_row

    # Header is right after intra title. Data of intra = intra_header_row+1 .. trans_title_row-2 (μείον kενά)
    intra_data_start = intra_title_row + 2  # title row, header row, then data
    intra_data_end   = trans_title_row - 1  # μπορεί να υπάρχει κενή γραμμή πριν τον επόμενο τίτλο
    trans_data_start = trans_title_row + 2
    trans_data_end   = ws.max_row

    def collect(start, end):
        out = []
        for r in range(start, end + 1):
            row = rows[r - 1]
            # Skip empty rows
            if not any(c for c in row):
                continue
            # Skip header rows that re-list "Συνεργάτης" etc.
            if any(str(c or "") == "Συνεργάτης" for c in row):
                continue
            out.append(row)
        return out

    intra = collect(intra_data_start, intra_data_end)
    trans = collect(trans_data_start, trans_data_end)

    print("\n--- INTRA ---")
    for r in intra: print(r)
    print("\n--- TRANS ---")
    for r in trans: print(r)

    def codes(rows_):
        return [str(row[1]) for row in rows_]  # στήλη 2 = Κωδ.Πελάτη

    intra_codes = codes(intra)
    trans_codes = codes(trans)

    print("\nintra codes:", intra_codes)
    print("trans codes:", trans_codes)

    errors = []
    # Expected:
    #   100 (καθαρό intra), 200 (μόνο ΤΡ) → intra
    #   300, 400 (μόνο TRANS), 500 (ΤΡ+TRANS), 600 (Bazaar), 700 (Bazaar) → trans
    for c in ("100", "200"):
        if c not in intra_codes:
            errors.append(f"Code {c} δεν είναι στο intra (βρέθηκε σε trans={c in trans_codes})")
    for c in ("300", "400", "500", "600", "700"):
        if c not in trans_codes:
            errors.append(f"Code {c} δεν είναι στο 'Με Μεταφορική' (βρέθηκε σε intra={c in intra_codes})")

    # Verify Πελάτες Ημέρας: το 500 (ΤΡ+TRANS) πρέπει να εμφανίζεται στο φύλλο
    # (στο "tr" bucket γιατί έχει ΤΡ flag — αυτό είναι ΑΛΛΟ contract, δεν το αλλάζουμε)
    ws_p = wb["Πελατες Ημερας"]
    p_rows = [[ws_p.cell(r, c).value for c in range(1, ws_p.max_column + 1)] for r in range(1, ws_p.max_row + 1)]
    p_codes = [str(row[1]) for row in p_rows[2:] if row[1] not in (None, "")]
    print("\nΠελατες Ημερας codes:", p_codes)
    for c in ("100", "200", "300", "400", "500", "600", "700"):
        if c not in p_codes:
            errors.append(f"Code {c} λείπει από 'Πελατες Ημερας' (regression)")

    if errors:
        print("\nFAIL:")
        for e in errors:
            print(" -", e)
        sys.exit(1)
    print("\nPASS")


if __name__ == "__main__":
    run()
