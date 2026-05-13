import os, glob
from openpyxl import load_workbook
folder = r'C:\Users\U0091\Desktop\Lista'
for path in glob.glob(os.path.join(folder, '*.xlsx')):
    print('='*80)
    print('FILE:', path)
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as e:
        print('  ERR:', e); continue
    for ws in wb.worksheets:
        print(f'  Sheet="{ws.title}" rows={ws.max_row} cols={ws.max_column}')
        for r in range(1, min(3, ws.max_row)+1):
            row = [ws.cell(r,c).value for c in range(1, ws.max_column+1)]
            print(f'    R{r}:', row)
